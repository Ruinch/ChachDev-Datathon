"""Read human notes first; retain AI suggestions as unreviewed, private provenance.

Never called by an inference module. Source workbooks are read, never modified.
Blank human notes do not mean a negative label; AI intervals are not ground truth.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from ml_tools.dataset import canonical_id, parse_timestamp


def digest(path: Path) -> str:
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def reconcile(manifest: dict, human_path: Path, ai_path: Path) -> dict:
    from openpyxl import load_workbook

    sources = {'human': {'file': human_path.name, 'sha256': digest(human_path)},
               'ai': {'file': ai_path.name, 'sha256': digest(ai_path)}}
    primary = {}
    wb = load_workbook(human_path, read_only=True, data_only=True)
    try:
        sheet = wb['labels']
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if row[0] is None:
                continue
            key = canonical_id(str(row[0]))
            if key in primary:
                raise ValueError('Duplicate normalized ID in primary workbook')
            name, point, label, *notes = row
            if label not in {'нарушение', 'без нарушения'}:
                raise ValueError('Unknown primary label')
            primary[key] = {'label': 'violation' if label == 'нарушение' else 'no_violation',
                'point_sec': parse_timestamp(point), 'note': str(notes[0] or '').strip() if notes else '',
                'source': {'role': 'human', 'sheet': sheet.title, 'row': row_number}}
    finally:
        wb.close()
    secondary = {}
    ids = {r['id'] for r in manifest['records']}
    def ai_key(filename):
        key = canonical_id(str(filename))
        # AI workbook used a local copy suffix for one source. Only resolve when
        # the exact ID is absent and the unsuffixed ID exists in this manifest.
        base = re.sub(r'\(\d+\)$', '', key)
        return base if key not in ids and base in ids else key
    wb = load_workbook(ai_path, read_only=True, data_only=True)
    try:
        sheet = wb['Реестр видео']
        for row_number, row in enumerate(sheet.iter_rows(min_row=6, values_only=True), 6):
            if len(row) < 8 or not row[1] or not str(row[1]).lower().endswith('.mov'):
                continue
            key = ai_key(row[1])
            if key in secondary:
                raise ValueError('Duplicate normalized ID in AI registry')
            secondary[key] = {'category': row[5], 'suggested_class': row[6], 'summary': row[7],
                'source': {'role': 'ai', 'sheet': sheet.title, 'row': row_number}, 'proposed_events': []}
        sheet = wb['События и гипотезы']
        for row_number, row in enumerate(sheet.iter_rows(min_row=6, values_only=True), 6):
            if len(row) < 13 or not row[1] or not str(row[1]).lower().endswith('.mov'):
                continue
            key = ai_key(row[1])
            event = {'rule_id': row[4], 'category': row[3], 'observation': row[6],
                'limitations': row[7], 'context_start_sec': row[9], 'context_end_sec': row[10],
                'start_sec': row[11], 'end_sec': row[12], 'human_verified': False,
                'source': {'role': 'ai', 'sheet': sheet.title, 'row': row_number}}
            if key in secondary:
                secondary[key]['proposed_events'].append(event)
    finally:
        wb.close()
    if set(primary) != ids or set(secondary) - ids:
        raise ValueError('Workbooks and manifest do not describe the same dataset')
    records = []
    for video in manifest['records']:
        human = primary[video['id']]
        if human['point_sec'] is not None and not 0 <= human['point_sec'] < video['duration_sec']:
            raise ValueError('Human timestamp outside source duration')
        ai = secondary.get(video['id'], {})
        informative = human['note'].casefold() not in {'', 'хз', 'не знаю', '?'}
        records.append({'id': video['id'], 'sha256': video['sha256'], 'duration_sec': video['duration_sec'],
            'split': video['split'], 'human': human, 'ai': ai,
            'preferred_observation': human['note'] if informative else ai.get('summary', ''),
            'preferred_observation_source': 'human' if informative else 'ai_unverified',
            'review_required': True})
    duplicates = {}
    for row in records:
        duplicates.setdefault(row['sha256'], []).append(row)
    groups = [{'ids': [r['id'] for r in rows],
               'human_labels': sorted({r['human']['label'] for r in rows}),
               'human_point_seconds': sorted({r['human']['point_sec'] for r in rows if r['human']['point_sec'] is not None}),
               'point_conflict': len({r['human']['point_sec'] for r in rows}) > 1}
              for rows in duplicates.values() if len(rows) > 1]
    return {'version': 1, 'sources': sources, 'policy': 'human_notes_first; AI intervals remain proposals',
        'records': records, 'duplicate_groups': groups,
        'summary': {'videos': len(records), 'unique_videos': len(duplicates),
                    'human_notes': sum(bool(r['human']['note']) for r in records),
                    'informative_human_notes': sum(r['preferred_observation_source'] == 'human' for r in records),
                    'human_verified_intervals': 0}}


def review_seed(reconciled: dict) -> dict:
    annotations = {}
    for row in reconciled['records']:
        human, ai = row['human'], row['ai']
        notes = 'Основной источник — заметка человека: ' + (human['note'] or 'нет описания')
        notes += f"\nИсходная метка: {human['label']}; ориентир: {human['point_sec']} с. Это точка, не границы события."
        if ai:
            notes += '\nВторичный разбор ИИ (не подтверждён человеком): ' + str(ai.get('summary', ''))
            for event in ai.get('proposed_events', []):
                notes += f"\nПредложение ИИ: {event['rule_id']}, {event['start_sec']}–{event['end_sec']} с; {event['observation']}"
        annotations[row['id']] = {'sha256': row['sha256'], 'reviewed': False,
            'source_group': 'organizer-single-camera', 'viewpoint': '', 'verdict': 'uncertain',
            'events': [], 'notes': notes}
    return {'version': 1, 'annotations': annotations}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--human', required=True, type=Path)
    parser.add_argument('--ai', required=True, type=Path)
    parser.add_argument('--manifest', type=Path, default=Path('.local_ml/dataset.json'))
    parser.add_argument('--output', type=Path, default=Path('.local_ml/reconciled'))
    args = parser.parse_args()
    if args.output.exists() and any(args.output.iterdir()):
        raise SystemExit('Use a new output directory to preserve existing reviews')
    reconciled = reconcile(json.loads(args.manifest.read_text(encoding='utf-8')), args.human, args.ai)
    args.output.mkdir(parents=True, exist_ok=True)
    for name, obj in [('sources.json', reconciled), ('review-seed.json', review_seed(reconciled))]:
        (args.output / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(reconciled['summary']))


if __name__ == '__main__':
    main()
